"""Núcleo: verificador, lote, rastreio, sitemap, entradas e exportação."""

import json
import sys
import tempfile
import threading
import unittest
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from servidor_teste import Manipulador, Servidor  # noqa: E402

from validaurl import classificar as C  # noqa: E402
from validaurl.entradas import urls_do_arquivo, urls_do_texto  # noqa: E402
from validaurl.exportar import exportar  # noqa: E402
from validaurl.links import extrair_links, mesmo_site, normalizar  # noqa: E402
from validaurl.lote import verificar_lote  # noqa: E402
from validaurl.rastreador import OpcoesRastreio, rastrear  # noqa: E402
from validaurl.sitemap import ErroSitemap, interpretar, urls_do_sitemap  # noqa: E402
from validaurl.verificador import Opcoes, novo_cliente, verificar  # noqa: E402


class TesteVerificador(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.srv = Servidor().__enter__()
        cls.opc = Opcoes(timeout=1.5)
        cls.cli = novo_cliente(cls.opc)

    @classmethod
    def tearDownClass(cls):
        cls.cli.close()
        cls.srv.__exit__()

    def checar(self, caminho, **kw):
        opc = Opcoes(**{"timeout": 1.5, **kw})
        with novo_cliente(opc) as cli:
            return verificar(self.srv.url(caminho), cli, opc)[0]

    def test_200(self):
        r = self.checar("/ok")
        self.assertEqual((r.status, r.categoria, r.n_redirects, r.metodo), (200, C.OK, 0, "HEAD"))

    def test_404(self):
        r = self.checar("/nada")
        self.assertEqual((r.status, r.categoria), (404, C.ERRO_CLIENTE))
        self.assertEqual(r.erro, "Not Found")
        self.assertEqual(r.metodo, "GET")  # HEAD >= 400 repete com GET

    def test_500(self):
        self.assertEqual(self.checar("/500").categoria, C.ERRO_SERVIDOR)

    def test_cadeia_redirect(self):
        r = self.checar("/redir2")
        self.assertEqual(r.status, 200)
        self.assertEqual(r.categoria, C.REDIRECIONADO)
        self.assertEqual([s for s, _ in r.redirects], [302, 301])
        self.assertEqual(r.status_inicial, 302)
        self.assertTrue(r.url_final.endswith("/ok"))

    def test_redirect_para_404(self):
        r = self.checar("/redir-404")
        self.assertEqual((r.status, r.categoria, r.n_redirects), (404, C.ERRO_CLIENTE, 1))

    def test_redirect_nao_seguido(self):
        r = self.checar("/redir", seguir_redirects=False)
        self.assertEqual((r.status, r.categoria), (301, C.REDIRECT_NAO_SEGUIDO))

    def test_loop(self):
        r = self.checar("/loop")
        self.assertIsNone(r.status)
        self.assertEqual(r.categoria, C.FALHA)
        self.assertIn("Loop", r.erro)

    def test_head_405_cai_para_get(self):
        r = self.checar("/head405")
        self.assertEqual((r.status, r.metodo), (200, "GET"))

    def test_head_mentiroso(self):
        self.assertEqual(self.checar("/head-mente").status, 200)

    def test_timeout(self):
        r = self.checar("/lento", timeout=0.5)
        self.assertEqual((r.categoria, r.erro), (C.FALHA, "Tempo esgotado"))
        self.assertLess(r.tempo_ms, 2500)  # não repete o GET depois do timeout no HEAD

    def test_conexao_recusada(self):
        opc = Opcoes(timeout=3)
        with novo_cliente(opc) as cli:
            r = verificar("http://127.0.0.1:1/", cli, opc)[0]
        self.assertEqual(r.categoria, C.FALHA)
        self.assertTrue(r.erro)

    def test_dns(self):
        opc = Opcoes(timeout=5)
        with novo_cliente(opc) as cli:
            r = verificar("https://nao-existe-validaurl.invalid/", cli, opc)[0]
        self.assertEqual(r.categoria, C.FALHA)
        self.assertEqual(r.erro, "Domínio não encontrado (DNS)")

    def test_url_invalida(self):
        with novo_cliente(self.opc) as cli:
            for u in ("ftp://x.com/a", "http://", "nada"):
                r = verificar(u, cli, self.opc)[0]
                self.assertEqual(r.categoria, C.FALHA, u)
                self.assertTrue(r.erro, u)

    def test_ler_html(self):
        opc = Opcoes(timeout=2)
        with novo_cliente(opc) as cli:
            r, html = verificar(self.srv.url("/site/"), cli, opc, ler_html=True)
            self.assertIn('href="a"', html)
            _, binario = verificar(self.srv.url("/img.png"), cli, opc, ler_html=True)
            self.assertIsNone(binario)


class TesteLote(unittest.TestCase):
    def test_lote_ordem_e_callback(self):
        with Servidor() as srv:
            urls = [srv.url(p) for p in ("/ok", "/nada", "/redir", "/500", "/ok?x=1")]
            vistos = []
            res = verificar_lote(urls, Opcoes(timeout=2, concorrencia=4, por_host=2), ao_resultado=vistos.append)
        self.assertEqual([r.url for r in res], urls)
        self.assertEqual(len(vistos), 5)
        self.assertEqual([r.status for r in res], [200, 404, 200, 500, 200])

    def test_cancelar(self):
        with Servidor() as srv:
            cancelar = threading.Event()
            cancelar.set()
            res = verificar_lote([srv.url("/ok")] * 20, Opcoes(timeout=2), cancelar=cancelar)
        self.assertEqual(res, [])


class TesteLinksEntradas(unittest.TestCase):
    def test_extrair(self):
        html = ('<base href="https://x.com/dir/"><a href="p1">1</a><a href="/p2#f">2</a><a href="p1#z">dup</a>'
                '<a href="mailto:a@b.c">m</a><a href="javascript:x()">j</a><img src="i.png" srcset="i2.png 2x, i3.png 3x">'
                '<link rel="preconnect" href="https://cdn.x.com"><link rel="stylesheet" href="s.css">'
                '<a href="ftp://x.com/f">ftp</a><a href="#topo">topo</a>')
        self.assertEqual(extrair_links(html, "https://x.com/"), [
            "https://x.com/dir/p1", "https://x.com/p2", "https://x.com/dir/i.png", "https://x.com/dir/i2.png",
            "https://x.com/dir/i3.png", "https://x.com/dir/s.css"])

    def test_normalizar(self):
        self.assertEqual(normalizar("HTTPS://X.com:443?b=2&a=1#f"), "https://x.com/?a=1&b=2")
        self.assertEqual(normalizar("http://x.com:8080/a"), "http://x.com:8080/a")

    def test_mesmo_site(self):
        self.assertTrue(mesmo_site("https://www.x.com/a", "https://x.com"))
        self.assertFalse(mesmo_site("https://blog.x.com/a", "https://x.com"))
        self.assertTrue(mesmo_site("https://blog.x.com/a", "https://x.com", subdominios=True))
        self.assertFalse(mesmo_site("https://xx.com/a", "https://x.com", subdominios=True))

    def test_texto(self):
        texto = ("veja https://a.com/x, e (https://b.com/y).\n"
                 "www.c.com.br/pagina\n"
                 "d.org\n"
                 "email@e.com\n"
                 "https://a.com/x\n"
                 "https://pt.wikipedia.org/wiki/Java_(linguagem)\n"
                 "texto qualquer sem link\n")
        self.assertEqual(urls_do_texto(texto), [
            "https://a.com/x", "https://b.com/y", "https://www.c.com.br/pagina", "https://d.org",
            "https://pt.wikipedia.org/wiki/Java_(linguagem)"])

    def test_arquivos(self):
        with tempfile.TemporaryDirectory() as d:
            d = Path(d)
            (d / "l.txt").write_text("https://a.com\nhttps://b.com\n", encoding="utf-8")
            (d / "l.csv").write_text("nome;site\nA;https://a.com/á\nB;b.com\n", encoding="cp1252")
            self.assertEqual(urls_do_arquivo(d / "l.txt"), ["https://a.com", "https://b.com"])
            self.assertEqual(urls_do_arquivo(d / "l.csv"), ["https://a.com/á", "https://b.com"])
            from openpyxl import Workbook
            wb = Workbook()
            ws = wb.active
            ws.append(["Site", "Obs"])
            ws.append(["https://x.com/1", "nada"])
            ws["A3"] = "clique aqui"
            ws["A3"].hyperlink = "https://y.com/2"
            wb.save(d / "l.xlsx")
            self.assertEqual(urls_do_arquivo(d / "l.xlsx"), ["https://x.com/1", "https://y.com/2"])


class TesteRastreio(unittest.TestCase):
    def test_rastreio_completo(self):
        Manipulador.contagem.clear()
        with Servidor() as srv:
            avisos = []
            res = rastrear(srv.url("/site/"), Opcoes(timeout=3, concorrencia=6),
                           OpcoesRastreio(respeitar_robots=True), avisos=avisos)
        por_url = {r.url.replace(srv.base, ""): r for r in res}
        for caminho in ("/site/", "/site/a", "/site/b", "/site/c", "/img.png", "/500", "/site/fundo/3",
                        "/site/privado"):
            self.assertIn(caminho, por_url)
        self.assertEqual(por_url["/site/c"].status, 404)
        self.assertEqual(por_url["/site/c"].encontrado_em, srv.url("/site/a"))
        self.assertEqual(por_url["/500"].encontrado_em, srv.url("/site/b"))
        # robots.txt: o privado é checado, mas não aberto
        self.assertNotIn("/site/so-pelo-privado", por_url)
        # externo checado (porta fechada) e não rastreado
        self.assertTrue(any("127.0.0.2" in u and r.categoria == C.FALHA for u, r in por_url.items()))
        # sem duplicatas (o "#topo" é o mesmo /site/b)
        self.assertEqual(len(res), len({r.url for r in res}))
        self.assertFalse(any(u.startswith("mailto") or "javascript" in u for u in por_url))

    def test_inicio_redireciona_para_fora(self):
        # A página inicial sai do site: não há links para seguir, e o programa avisa o porquê.
        with Servidor() as srv:
            avisos = []
            rastrear(srv.url("/para-fora"), Opcoes(timeout=2), OpcoesRastreio(checar_externos=False), avisos=avisos)
        self.assertEqual(len(avisos), 1)
        self.assertIn("redireciona para localhost", avisos[0])
        self.assertIn("aba Lista de URLs", avisos[0])
        with Servidor() as srv:
            avisos = []
            rastrear(srv.url("/ok"), Opcoes(timeout=2), OpcoesRastreio(respeitar_robots=False), avisos=avisos)
        self.assertEqual(avisos, [], "sem redirect para fora, nenhum aviso")

    def test_limites(self):
        with Servidor() as srv:
            res = rastrear(srv.url("/site/"), Opcoes(timeout=3),
                           OpcoesRastreio(profundidade=1, checar_externos=False, respeitar_robots=False))
            caminhos = {r.url.replace(srv.base, "") for r in res}
            self.assertIn("/site/fundo/1", caminhos)      # achado na profundidade 2, checado
            self.assertNotIn("/site/fundo/2", caminhos)   # mas não aberto
            self.assertIn("/site/so-pelo-privado", caminhos)  # sem robots
            self.assertFalse(any("127.0.0.2" in c for c in caminhos))

            avisos = []
            res = rastrear(srv.url("/site/"), Opcoes(timeout=3),
                           OpcoesRastreio(max_paginas=1, checar_externos=False), avisos=avisos)
            self.assertNotIn("/site/c", {r.url.replace(srv.base, "") for r in res})
            self.assertTrue(avisos)

    def test_inicio_invalido(self):
        with self.assertRaises(ValueError):
            rastrear("ftp://x")


class TesteSitemap(unittest.TestCase):
    def test_sitemap_e_index(self):
        with Servidor() as srv, novo_cliente(Opcoes(timeout=3)) as cli:
            simples = urls_do_sitemap(cli, srv.url("/sitemap.xml"))
            self.assertEqual(simples, [srv.url("/ok"), srv.url("/500"), srv.url("/redir")])
            avisos = []
            todas = urls_do_sitemap(cli, srv.url("/sitemap_index.xml"), avisos)
            self.assertEqual(todas, [srv.url(p) for p in ("/ok", "/500", "/redir", "/head405")])
            self.assertEqual(len(avisos), 1)  # o sitemap-sumiu
            # só o site: descobre pelo robots.txt
            self.assertEqual(urls_do_sitemap(cli, srv.url("/")), todas)
            with self.assertRaises(ErroSitemap):
                urls_do_sitemap(cli, srv.url("/sitemap-bomba.xml"))
            with self.assertRaises(ErroSitemap):
                urls_do_sitemap(cli, srv.url("/sitemap-nao-tem.xml"))

    def test_interpretar_sem_namespace(self):
        self.assertEqual(interpretar(b"<urlset><url><loc>https://a.com</loc></url></urlset>"), (["https://a.com"], []))


class TesteExportar(unittest.TestCase):
    def test_formatos(self):
        with Servidor() as srv:
            res = verificar_lote([srv.url("/ok"), srv.url("/redir2"), srv.url("/nada")], Opcoes(timeout=2))
        with tempfile.TemporaryDirectory() as d:
            for ext in ("csv", "xlsx", "json"):
                exportar(res, Path(d) / f"r.{ext}")
            csv_txt = (Path(d) / "r.csv").read_text(encoding="utf-8-sig")
            self.assertTrue(csv_txt.startswith("URL;Status;Categoria;Tipo;"))
            self.assertIn(";HTML;", csv_txt, "o tipo do arquivo entra no CSV, com nome legível")
            self.assertIn("302 ", csv_txt)
            dados = json.loads((Path(d) / "r.json").read_text(encoding="utf-8"))
            self.assertEqual(len(dados["resultados"]), 3)
            from openpyxl import load_workbook
            wb = load_workbook(Path(d) / "r.xlsx")
            ws = wb.active
            self.assertEqual(ws.max_row, 4)
            self.assertEqual(ws["C4"].value, C.ERRO_CLIENTE)
            self.assertEqual(ws["D2"].value, "HTML", "coluna Tipo na planilha")
            wb.close()


class TesteRotuloTipo(unittest.TestCase):
    def test_nomes_legiveis(self):
        from validaurl.classificar import rotulo_tipo
        self.assertEqual(rotulo_tipo("text/html"), "HTML")
        self.assertEqual(rotulo_tipo("application/pdf"), "PDF")
        self.assertEqual(rotulo_tipo("image/png"), "PNG")
        self.assertEqual(rotulo_tipo("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "Excel (.xlsx)")
        self.assertEqual(rotulo_tipo("TEXT/HTML"), "HTML", "sem diferenciar maiúsculas")

    def test_desconhecido_e_vazio(self):
        from validaurl.classificar import rotulo_tipo
        self.assertEqual(rotulo_tipo(""), "", "servidor que não informou o tipo fica vazio")
        self.assertEqual(rotulo_tipo("application/x-algo-novo"), "application/x-algo-novo", "desconhecido mostra o original")


if __name__ == "__main__":
    unittest.main()
