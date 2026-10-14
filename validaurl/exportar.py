"""Exporta resultados em CSV, XLSX ou JSON."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from . import VERSAO
from .classificar import CORES
from .verificador import Resultado

COLUNAS = ["URL", "Status", "Categoria", "Tipo", "URL final", "Redirects", "Tempo (ms)", "Erro", "Encontrado em"]


def linha(r: Resultado) -> list:
    return [
        r.url,
        r.status if r.status is not None else "",
        r.categoria,
        r.tipo,
        r.url_final if r.url_final and r.url_final != r.url else "",
        " → ".join(f"{s} {u}" for s, u in r.redirects),
        r.tempo_ms,
        r.erro,
        r.encontrado_em,
    ]


def para_csv(resultados: list[Resultado], caminho: str | Path) -> None:
    # ";" e BOM: o Excel em português abre direto, sem assistente de importação.
    with open(caminho, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUNAS)
        for r in resultados:
            w.writerow(linha(r))


def para_json(resultados: list[Resultado], caminho: str | Path) -> None:
    dados = {"gerado_por": f"ValidaURL {VERSAO}", "resultados": [r.como_dict() for r in resultados]}
    Path(caminho).write_text(json.dumps(dados, ensure_ascii=False, indent=2), encoding="utf-8")


def para_xlsx(resultados: list[Resultado], caminho: str | Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    livro = Workbook()
    aba = livro.active
    aba.title = "Resultados"
    aba.append(COLUNAS)
    for c in aba[1]:
        c.font = Font(bold=True)
    preenchimentos = {k: PatternFill("solid", fgColor=v.lstrip("#")) for k, v in CORES.items()}
    for r in resultados:
        aba.append(linha(r))
        cor = preenchimentos.get(r.categoria)
        if cor:
            for c in aba[aba.max_row]:
                c.fill = cor
    for letra, largura in zip("ABCDEFGHI", (60, 8, 20, 24, 50, 40, 11, 40, 50)):
        aba.column_dimensions[letra].width = largura
    aba.freeze_panes = "A2"
    aba.auto_filter.ref = aba.dimensions
    livro.save(caminho)


def exportar(resultados: list[Resultado], caminho: str | Path) -> None:
    ext = Path(caminho).suffix.lower()
    if ext == ".xlsx":
        para_xlsx(resultados, caminho)
    elif ext == ".json":
        para_json(resultados, caminho)
    else:
        para_csv(resultados, caminho)
