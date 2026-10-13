"""Lê URLs de texto colado, TXT, CSV e XLSX."""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from .links import normalizar

RE_URL = re.compile(r"https?://[^\s\"'<>\]\[{}|\\^`]+", re.IGNORECASE)
# Sem esquema: "www.site.com.br/x" ou "site.com.br" sozinho na célula/linha.
RE_DOMINIO = re.compile(r"^(?:www\.)?[a-z0-9-]+(?:\.[a-z0-9-]+)*\.[a-z]{2,}(?::\d+)?(?:[/?#]\S*)?$", re.IGNORECASE)
PONTUACAO_FINAL = ".,;:!?)'\"»"


def _limpar(url: str) -> str:
    url = url.strip()
    # Parêntese de fechamento só sai se não houver um de abertura (links da Wikipédia).
    while url and url[-1] in PONTUACAO_FINAL:
        if url[-1] == ")" and url.count("(") >= url.count(")"):
            break
        url = url[:-1]
    return url


def urls_do_texto(texto: str) -> list[str]:
    """Todas as URLs de um texto livre, na ordem, sem duplicatas."""
    saida: list[str] = []
    for linha in texto.splitlines():
        achadas = RE_URL.findall(linha)
        if achadas:
            saida.extend(_limpar(u) for u in achadas)
            continue
        for pedaco in re.split(r"[\s;,\t]+", linha.strip()):
            pedaco = _limpar(pedaco)
            if pedaco and "@" not in pedaco and RE_DOMINIO.match(pedaco):
                saida.append("https://" + pedaco)
    return deduplicar(saida)


def deduplicar(urls) -> list[str]:
    vistos, saida = set(), []
    for u in urls:
        if not u:
            continue
        chave = normalizar(u)
        if chave not in vistos:
            vistos.add(chave)
            saida.append(u)
    return saida


def _ler_texto(caminho: Path) -> str:
    bruto = caminho.read_bytes()
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return bruto.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return bruto.decode("utf-8", errors="replace")


def urls_do_csv(texto: str) -> list[str]:
    try:
        dialeto = csv.Sniffer().sniff(texto[:4096], delimiters=";,\t|")
    except csv.Error:
        dialeto = csv.excel
    celulas = []
    for linha in csv.reader(io.StringIO(texto), dialeto):
        celulas.extend(linha)
    return urls_do_texto("\n".join(celulas))


def urls_do_xlsx(caminho: Path) -> list[str]:
    from openpyxl import load_workbook

    livro = load_workbook(caminho, data_only=True)
    pedacos: list[str] = []
    try:
        for aba in livro.worksheets:
            for linha in aba.iter_rows():
                for celula in linha:
                    if celula.hyperlink is not None and celula.hyperlink.target:
                        pedacos.append(celula.hyperlink.target)
                    if isinstance(celula.value, str):
                        pedacos.append(celula.value)
    finally:
        livro.close()
    return urls_do_texto("\n".join(pedacos))


def urls_do_arquivo(caminho: str | Path) -> list[str]:
    caminho = Path(caminho)
    ext = caminho.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        return urls_do_xlsx(caminho)
    texto = _ler_texto(caminho)
    if ext in (".csv", ".tsv"):
        return urls_do_csv(texto)
    if ext in (".html", ".htm"):
        from .links import extrair_links
        return deduplicar(extrair_links(texto, caminho.resolve().as_uri()))
    return urls_do_texto(texto)
